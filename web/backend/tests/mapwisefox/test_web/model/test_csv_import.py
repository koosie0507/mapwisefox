import pytest
from mapwisefox.web.model import WorkbookValidationError, convert_csv_to_workbook
from openpyxl import Workbook, load_workbook


WOS_CSV = (
    b"Author Full Names,Article Title,Source Title,Abstract,"
    b"Author Keywords,Publication Year,DOI,DOI Link\n"
    b"John Doe,Test Title,Test Source,Test Abstract,Key1;Key2,"
    b"2023,10.1000/1,http://doi.org/1"
)

SCOPUS_CSV = (
    b"Author names,Title,Year,Source title,DOI,Link,Abstract,Index Keywords\n"
    b"Jane Roe,Scopus Title,2024,Scopus Source,10.2000/2,http://link/2,Scopus Abstract,Scopus;Keys"
)


@pytest.fixture
def explicit_field_mappings():
    return {"title": "My Title", "authors": "My Authors"}


def test_convert_materializes_workbook_with_records(tmp_path):
    converted, _ = convert_csv_to_workbook(WOS_CSV, tmp_path)

    worksheet = load_workbook(converted).active
    assert worksheet.cell(1, 1).value == "Author Full Names"
    assert worksheet.cell(2, 2).value == "Test Title"
    assert worksheet.cell(2, 6).value == "2023"
    assert worksheet.cell(2, 7).value == "10.1000/1"


def test_convert_keeps_values_as_strings(tmp_path):
    converted, _ = convert_csv_to_workbook(b"title,year\nT1,2025\n", tmp_path)

    worksheet = load_workbook(converted).active
    assert worksheet.cell(2, 2).value == "2025"


def test_convert_preserves_blank_record_rows(tmp_path):
    converted, _ = convert_csv_to_workbook(
        b"title,year\nT1,2025\n\nT2,2024\n", tmp_path
    )

    worksheet = load_workbook(converted).active
    assert worksheet.cell(3, 1).value in (None, "")
    assert worksheet.cell(4, 1).value == "T2"


def test_convert_pads_ragged_records(tmp_path):
    converted, _ = convert_csv_to_workbook(b"title,authors\nT1,A1\nT2\n", tmp_path)

    worksheet = load_workbook(converted).active
    assert worksheet.cell(3, 1).value == "T2"
    assert worksheet.cell(3, 2).value in (None, "")


def test_convert_resolves_wos_mappings(tmp_path):
    converted, mappings = convert_csv_to_workbook(WOS_CSV, tmp_path)

    assert mappings["title"] == "Article Title"
    assert mappings["authors"] == "Author Full Names"
    assert mappings["publicationDate"] == "Publication Year"
    assert mappings["url"] == "DOI Link"


def test_convert_resolves_scopus_mappings(tmp_path):
    converted, mappings = convert_csv_to_workbook(SCOPUS_CSV, tmp_path)

    assert mappings["authors"] == "Author names"
    assert mappings["publicationDate"] == "Year"
    assert mappings["publicationVenue"] == "Source title"
    assert mappings["keywords"] == "Index Keywords"


def test_convert_keeps_first_column_per_field(tmp_path):
    content = b"Article Title,Source Title,Abbreviated Source Title\nT,S,S2\n"

    converted, mappings = convert_csv_to_workbook(content, tmp_path)

    assert mappings["title"] == "Article Title"
    assert mappings["publicationVenue"] == "Source Title"
    assert "Abbreviated Source Title" not in mappings.values()


def test_convert_returns_explicit_mappings(tmp_path, explicit_field_mappings):
    converted, mappings = convert_csv_to_workbook(
        b"My Title,My Authors\nT,A\n", tmp_path, explicit_field_mappings
    )

    assert mappings == {"title": "My Title", "authors": "My Authors"}


def test_convert_resolves_nothing_without_usable_headers(tmp_path):
    converted, mappings = convert_csv_to_workbook(b",\n,\n", tmp_path)

    assert mappings == {}
    worksheet = load_workbook(converted).active
    assert worksheet.cell(1, 1).value in (None, "")
    assert worksheet.cell(2, 1).value in (None, "")


def test_convert_cleans_up_when_saving_fails(tmp_path, monkeypatch):
    def failing_save(self, filename):
        raise RuntimeError("disk full")

    monkeypatch.setattr(Workbook, "save", failing_save)

    with pytest.raises(RuntimeError, match="disk full"):
        convert_csv_to_workbook(b"title,authors\nT,A\n", tmp_path)

    assert list(tmp_path.glob("*.xlsx")) == []


def test_convert_tolerates_byte_order_mark(tmp_path):
    converted, mappings = convert_csv_to_workbook(
        b"\xef\xbb\xbftitle,authors\nT,A\n", tmp_path
    )

    worksheet = load_workbook(converted).active
    assert worksheet.cell(1, 1).value == "title"
    assert mappings == {"title": "title", "authors": "authors"}


def test_convert_rejects_empty_csv(tmp_path):
    with pytest.raises(WorkbookValidationError) as error:
        convert_csv_to_workbook(b"", tmp_path)

    assert error.value.code == "missing_header"


def test_convert_rejects_undecodable_csv(tmp_path):
    with pytest.raises(WorkbookValidationError) as error:
        convert_csv_to_workbook(b"\xff\xfe\x00binary", tmp_path)

    assert error.value.code == "invalid_workbook"
